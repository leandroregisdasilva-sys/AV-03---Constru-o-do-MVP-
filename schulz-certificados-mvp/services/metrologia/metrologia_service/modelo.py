"""Gera o modelo de planilha de entrada (formato simples do MVP)."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CABECALHO = ["Ponto", "Valor padrão (mm)", "LEA (mm)", "Leitura 1", "Leitura 2"]
EXEMPLO = [
    ("EXEMPLO_DIAMETRO_1", 10.000, 0.010, 10.0021, 10.0027),
    ("EXEMPLO_DIAMETRO_2", 25.000, 0.010, 24.9962, 24.9958),
    ("EXEMPLO_PLANICIDADE", 0.000, 0.050, 0.0123, 0.0131),
]


def gerar_modelo_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leituras"
    ws["A1"] = "Modelo de planilha – Análise metrológica (MVP Schulz)"
    ws["A1"].font = Font(bold=True, size=13)
    rotulos = [("Cliente:", "Setor solicitante"), ("Técnico:", "Nome do metrologista"),
               ("Código:", "Código do instrumento"), ("Denominação:", "Descrição do instrumento"),
               ("Data:", None)]
    for i, (rot, val) in enumerate(rotulos, start=3):
        ws.cell(row=i, column=1, value=rot).font = Font(bold=True)
        ws.cell(row=i, column=2, value=val)
    linha_cab = 9
    for col, titulo in enumerate(CABECALHO, start=1):
        c = ws.cell(row=linha_cab, column=col, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="151A2E")
        c.alignment = Alignment(horizontal="center")
    for r, linha in enumerate(EXEMPLO, start=linha_cab + 1):
        for col, valor in enumerate(linha, start=1):
            ws.cell(row=r, column=col, value=valor)
    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 24 if col == 1 else 18
    ws2 = wb.create_sheet("Instruções")
    texto = [
        "Como preencher",
        "1. Uma linha por ponto de ensaio, abaixo do cabeçalho.",
        "2. 'Valor padrão' é o valor de referência do padrão (valor nominal).",
        "3. 'LEA' é o Limite de Erro Admissível (simétrico). Alternativa: duas colunas 'Tol. +' e 'Tol. -'.",
        "4. Informe ao menos 2 leituras por ponto; pode haver 'Leitura 3', 'Leitura 4'…",
        "5. Use ponto ou vírgula como separador decimal. Valores em mm.",
        "6. Apague as linhas de exemplo antes de enviar.",
        "O relatório bruto do Calypso (Characteristic/Actual/Nominal/Upper Tol/Lower Tol) também é aceito.",
    ]
    for i, t in enumerate(texto, start=1):
        ws2.cell(row=i, column=1, value=t)
    ws2["A1"].font = Font(bold=True, size=13)
    ws2.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
